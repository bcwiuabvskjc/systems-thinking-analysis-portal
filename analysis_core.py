"""Core analysis functions for the Systems Thinking Analysis Portal."""
from __future__ import annotations

import io
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

KSIM_METHOD = (
    "Documented simplified KSIM-style variant. Variables are normalized to [0,1]. "
    "For each target, positive and negative exposures are calculated from signed "
    "interaction coefficients weighted by current source states. The exponent P "
    "is (negative exposure + epsilon)/(positive exposure + epsilon), and the update "
    "is x_next = x_current ** (P * dt). This is not guaranteed to match every "
    "published KSIM formulation or a course-specific variant."
)


def validate_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    if df.empty:
        rows.append({"Severity": "Error", "Check": "Empty dataset",
                     "Details": "The CSV has no data rows."})
    if len(df.columns) == 0:
        rows.append({"Severity": "Error", "Check": "No columns",
                     "Details": "No columns were detected."})
    duplicates = df.columns[df.columns.duplicated()].tolist()
    if duplicates:
        rows.append({"Severity": "Error", "Check": "Duplicate column names",
                     "Details": ", ".join(map(str, duplicates))})
    missing = int(df.isna().sum().sum())
    if missing:
        rows.append({"Severity": "Warning", "Check": "Missing values",
                     "Details": f"{missing} missing cells. Values were not imputed."})
    if not rows:
        rows.append({"Severity": "Pass", "Check": "Basic validation",
                     "Details": "No basic structural issues detected."})
    return pd.DataFrame(rows, columns=["Severity", "Check", "Details"])


def coerce_numeric(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    for col in result.columns:
        if result[col].dtype == object:
            values = result[col].astype(str).str.strip()
            nonblank = values.ne("") & values.ne("nan")
            parsed = pd.to_numeric(values.where(nonblank), errors="coerce")
            if nonblank.sum() > 0 and parsed[nonblank].notna().all():
                result[col] = parsed
    return result


def detect_cross_impact(df, source_col=None, target_col=None, value_col=None):
    if (source_col and target_col and value_col
            and all(c in df.columns for c in [source_col, target_col, value_col])):
        work = df[[source_col, target_col, value_col]].copy()
        work[value_col] = pd.to_numeric(work[value_col], errors="coerce")
        if work.isna().any().any():
            return None, "Mapped relationship rows contain missing or non-numeric values."
        if work.duplicated([source_col, target_col]).any():
            return None, "Duplicate source-target pairs found."
        matrix = work.pivot(index=source_col, columns=target_col, values=value_col)
        if (matrix.shape[0] != matrix.shape[1]
                or set(matrix.index.astype(str)) != set(matrix.columns.astype(str))):
            return None, "Long-form data must define a complete square set of source and target factors."
        matrix.index = matrix.index.astype(str)
        matrix.columns = matrix.columns.astype(str)
        matrix = matrix.reindex(index=sorted(matrix.index), columns=sorted(matrix.columns))
        return matrix, "Built from mapped relationship data."

    if df.shape[1] >= 2:
        labels = df.iloc[:, 0].astype(str)
        factors = [str(c) for c in df.columns[1:]]
        if len(factors) == len(labels) and set(labels) == set(factors):
            values = df.iloc[:, 1:].copy()
            values.columns = factors
            values.index = labels
            numeric = values.apply(pd.to_numeric, errors="coerce")
            if not numeric.isna().any().any():
                numeric = numeric.reindex(
                    index=sorted(numeric.index), columns=sorted(numeric.columns)
                )
                return numeric, "Detected square matrix."
    return None, "No unambiguous square matrix detected."


def cross_impact_metrics(matrix):
    numeric = matrix.apply(pd.to_numeric, errors="raise")
    metrics = pd.DataFrame({
        "Factor": numeric.index.astype(str),
        "Influence (row sum)": numeric.sum(axis=1).to_numpy(),
        "Dependence (column sum)": numeric.sum(axis=0).to_numpy(),
    })
    median_i = metrics["Influence (row sum)"].median()
    median_d = metrics["Dependence (column sum)"].median()
    metrics["Median-based descriptive group"] = metrics.apply(
        lambda row: (
            ("High" if row["Influence (row sum)"] >= median_i else "Low")
            + " influence / "
            + ("high" if row["Dependence (column sum)"] >= median_d else "low")
            + " dependence"
        ), axis=1
    )
    definitions = pd.DataFrame([
        {"Metric": "Influence", "Definition": "Row sum",
         "Interpretation": "Outgoing impact under the selected matrix convention."},
        {"Metric": "Dependence", "Definition": "Column sum",
         "Interpretation": "Incoming impact under the selected matrix convention."},
    ])
    return metrics, definitions


def run_ksim(initial_values, interaction_matrix, steps=20, dt=1.0, epsilon=1e-9):
    """Run a documented simplified KSIM-style simulation.

    Initial values must be normalized to [0,1].
    Matrix rows are target variables; columns are source variables.
    Positive coefficients enhance a target; negative coefficients inhibit it.
    """
    x0 = np.asarray(initial_values, dtype=float)
    matrix = np.asarray(interaction_matrix, dtype=float)
    if x0.ndim != 1 or matrix.shape != (len(x0), len(x0)):
        raise ValueError("Initial values must be a vector and the interaction matrix must be square with matching dimensions.")
    if not np.isfinite(x0).all() or not np.isfinite(matrix).all():
        raise ValueError("Inputs must contain only finite numeric values.")
    if ((x0 < 0) | (x0 > 1)).any():
        raise ValueError("Initial normalized values must be between 0 and 1.")
    if not 1 <= steps <= 500:
        raise ValueError("Choose between 1 and 500 simulation steps.")
    if not 0 < dt <= 5:
        raise ValueError("The time-step multiplier must be greater than 0 and at most 5.")

    x = np.clip(x0, epsilon, 1.0)
    history = [x.copy()]
    exponents = []
    for _ in range(int(steps)):
        next_x = x.copy()
        current_exponents = []
        for i in range(len(x)):
            weighted = matrix[i, :] * x
            positive = float(np.clip(weighted, 0, None).sum())
            negative = float(np.clip(-weighted, 0, None).sum())
            exponent = (negative + epsilon) / (positive + epsilon)
            exponent = float(np.clip(exponent, 0.05, 20.0))
            current_exponents.append(exponent)
            next_x[i] = np.clip(x[i] ** (exponent * dt), epsilon, 1.0)
        x = next_x
        history.append(x.copy())
        exponents.append(current_exponents)
    return np.asarray(history), np.asarray(exponents)


def denormalize_ksim(normalized, minimums, maximums):
    minimums = np.asarray(minimums, dtype=float)
    maximums = np.asarray(maximums, dtype=float)
    if (maximums <= minimums).any():
        raise ValueError("Each maximum must be greater than its minimum.")
    return minimums + np.asarray(normalized) * (maximums - minimums)


def build_excel(original, filename, cleaned, validation, params, matrix, metrics,
                numeric_summary, notes, ksim_results=None, ksim_exponents=None):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        original.to_excel(writer, sheet_name="Original Data", index=False)
        cleaned.to_excel(writer, sheet_name="Cleaned Data", index=False)
        validation.to_excel(writer, sheet_name="Data Validation", index=False)
        params.to_excel(writer, sheet_name="Parameters and Assumptions", index=False)
        if matrix is not None:
            matrix.rename_axis("Factor").to_excel(writer, sheet_name="Cross-Impact Matrix")
        if metrics is not None:
            metrics.to_excel(writer, sheet_name="Influence Dependence", index=False)
        if numeric_summary is not None and not numeric_summary.empty:
            numeric_summary.to_excel(writer, sheet_name="Numeric Summary", index=False)
        if ksim_results is not None:
            ksim_results.to_excel(writer, sheet_name="KSIM Results", index=False)
        if ksim_exponents is not None:
            ksim_exponents.to_excel(writer, sheet_name="KSIM Exponents", index=False)
        pd.DataFrame({"Notes and Interpretation": notes}).to_excel(
            writer, sheet_name="Executive Summary", index=False
        )

    buffer.seek(0)
    workbook = load_workbook(buffer)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column_cells in sheet.columns:
            width = max(
                (min(len(str(cell.value or "")), 60)
                 for cell in list(column_cells)[:250]),
                default=10,
            )
            sheet.column_dimensions[
                get_column_letter(column_cells[0].column)
            ].width = max(12, min(width + 2, 42))
        sheet.sheet_view.showGridLines = False

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
